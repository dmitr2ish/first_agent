from __future__ import annotations

import json
import os
import re
import sqlite3
import tempfile
import unicodedata
import warnings
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from project_config import REQUIRED_WORKFLOW_STEPS


WORKFLOW_TELEMETRY_KEY = "workflow_telemetry"

INTERACTION_MIGRATION_COLUMNS = {
    "tool_call_count": "INTEGER NOT NULL DEFAULT 0",
    "tool_names": "TEXT NOT NULL DEFAULT '[]'",
    "tool_error_count": "INTEGER NOT NULL DEFAULT 0",
    "evidence_read_performed": "INTEGER NOT NULL DEFAULT 0",
    "citation_count": "INTEGER NOT NULL DEFAULT 0",
    "used_source_ids": "TEXT NOT NULL DEFAULT '[]'",
    "invalid_source_ids": "TEXT NOT NULL DEFAULT '[]'",
    "source_count": "INTEGER NOT NULL DEFAULT 0",
    "researcher_success": "INTEGER",
    "validator_success": "INTEGER",
    "final_status": "TEXT NOT NULL DEFAULT 'unknown'",
    "final_safety_blocked": "INTEGER NOT NULL DEFAULT 0",
    "final_citation_count": "INTEGER NOT NULL DEFAULT 0",
    "final_used_source_ids": "TEXT NOT NULL DEFAULT '[]'",
    "final_invalid_source_ids": "TEXT NOT NULL DEFAULT '[]'",
    "final_output_policy_pass": "INTEGER",
    "final_safety_policy_pass": "INTEGER",
    "evidence_source_ids": "TEXT NOT NULL DEFAULT '[]'",
    "benchmark_scenario_id": "TEXT",
    "benchmark_scenario_type": "TEXT",
    "benchmark_turn_id": "TEXT",
    "benchmark_turn_number": "INTEGER",
    "benchmark_pass": "INTEGER",
    "benchmark_technical_pass": "INTEGER",
    "benchmark_semantic_pass": "INTEGER",
    "benchmark_release_pass": "INTEGER",
    "benchmark_failure_reasons": "TEXT NOT NULL DEFAULT '[]'",
    "semantic_failure_reasons": "TEXT NOT NULL DEFAULT '[]'",
}

TOOL_CALL_MIGRATION_COLUMNS = {
    "evidence_source_ids": "TEXT NOT NULL DEFAULT '[]'",
}

ILLEGAL_XLSX_CONTROL_RE = re.compile(
    r"[\x00-\x08\x0B-\x0C\x0E-\x1F]"
)


class ReportExportWarning(UserWarning):
    """SQLite was updated, but the derived XLSX export failed."""


@dataclass(frozen=True)
class ExperimentConfig:
    """
    Конфигурация одного экспериментального запуска.

    Новый run создаётся при запуске программы или по команде /new_run.
    Все вопросы до следующего /new_run относятся к одной конфигурации.
    """

    model_name: str
    model_options: dict[str, Any]
    workflow_version: str
    prompt_version: str
    source_collector_version: str
    max_source_pages: int
    allowed_domains: list[str]
    git_commit: str
    notes: str = ""
    benchmark_version: str = ""
    benchmark_config_digest: str = ""
    component_versions: dict[str, str] = field(default_factory=dict)


class ExperimentReporter:
    """
    SQLite — основной источник данных.
    XLSX — автоматически пересоздаваемый сравнительный отчёт.
    """

    def __init__(
        self,
        db_path: str = "experiment_results.db",
        xlsx_path: str = "visa_model_comparison.xlsx",
    ) -> None:
        self.db_path = Path(db_path)
        self.xlsx_path = Path(xlsx_path)
        self.last_export_error: str | None = None

        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        self.xlsx_path.parent.mkdir(parents=True, exist_ok=True)

        self._initialize_database()

    # =========================================================
    # DATABASE
    # =========================================================

    def _connect(self) -> sqlite3.Connection:
        connection = sqlite3.connect(self.db_path)
        connection.row_factory = sqlite3.Row

        # SQLite может хранить несколько читателей и одного писателя.
        connection.execute("PRAGMA journal_mode=WAL")
        connection.execute("PRAGMA foreign_keys=ON")

        return connection

    def _initialize_database(self) -> None:
        with self._connect() as connection:
            connection.executescript(
                """
                CREATE TABLE IF NOT EXISTS experiment_runs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    run_number INTEGER NOT NULL UNIQUE,
                    started_at_utc TEXT NOT NULL,
                    finished_at_utc TEXT,
                    model_name TEXT NOT NULL,
                    configuration_json TEXT NOT NULL,
                    notes TEXT NOT NULL DEFAULT ''
                );

                CREATE TABLE IF NOT EXISTS interactions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    experiment_run_id INTEGER NOT NULL,
                    question_number INTEGER NOT NULL,
                    workflow_run_id TEXT,
                    session_id TEXT,
                    created_at_utc TEXT NOT NULL,

                    question TEXT NOT NULL,
                    answer TEXT NOT NULL,

                    status TEXT NOT NULL,
                    duration_seconds REAL,
                    input_tokens INTEGER,
                    output_tokens INTEGER,
                    total_tokens INTEGER,
                    time_to_first_token REAL,

                    sources_json TEXT NOT NULL DEFAULT '[]',
                    step_outputs_json TEXT NOT NULL DEFAULT '{}',
                    metrics_json TEXT NOT NULL DEFAULT '{}',
                    tool_call_count INTEGER NOT NULL DEFAULT 0,
                    tool_names TEXT NOT NULL DEFAULT '[]',
                    tool_error_count INTEGER NOT NULL DEFAULT 0,
                    evidence_read_performed INTEGER NOT NULL DEFAULT 0,
                    citation_count INTEGER NOT NULL DEFAULT 0,
                    used_source_ids TEXT NOT NULL DEFAULT '[]',
                    invalid_source_ids TEXT NOT NULL DEFAULT '[]',
                    source_count INTEGER NOT NULL DEFAULT 0,
                    researcher_success INTEGER,
                    validator_success INTEGER,
                    final_status TEXT NOT NULL DEFAULT 'unknown',
                    final_safety_blocked INTEGER NOT NULL DEFAULT 0,
                    final_citation_count INTEGER NOT NULL DEFAULT 0,
                    final_used_source_ids TEXT NOT NULL DEFAULT '[]',
                    final_invalid_source_ids TEXT NOT NULL DEFAULT '[]',
                    final_output_policy_pass INTEGER,
                    final_safety_policy_pass INTEGER,
                    evidence_source_ids TEXT NOT NULL DEFAULT '[]',
                    benchmark_scenario_id TEXT,
                    benchmark_scenario_type TEXT,
                    benchmark_turn_id TEXT,
                    benchmark_turn_number INTEGER,
                    benchmark_pass INTEGER,
                    benchmark_technical_pass INTEGER,
                    benchmark_semantic_pass INTEGER,
                    benchmark_release_pass INTEGER,
                    benchmark_failure_reasons TEXT NOT NULL DEFAULT '[]',
                    semantic_failure_reasons TEXT NOT NULL DEFAULT '[]',
                    error TEXT,

                    FOREIGN KEY (experiment_run_id)
                        REFERENCES experiment_runs(id)
                        ON DELETE CASCADE,

                    UNIQUE (experiment_run_id, question_number)
                );

                CREATE INDEX IF NOT EXISTS idx_interactions_run
                    ON interactions(experiment_run_id);

                CREATE INDEX IF NOT EXISTS idx_interactions_workflow_run
                    ON interactions(workflow_run_id);

                CREATE TABLE IF NOT EXISTS tool_calls (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    interaction_id INTEGER NOT NULL,
                    sequence_number INTEGER NOT NULL,
                    step_name TEXT NOT NULL,
                    agent_name TEXT NOT NULL,
                    attempt INTEGER NOT NULL DEFAULT 1,
                    tool_name TEXT NOT NULL,
                    arguments_json TEXT NOT NULL DEFAULT '{}',
                    success INTEGER NOT NULL,
                    error INTEGER NOT NULL,
                    duration_seconds REAL,
                    evidence_read INTEGER NOT NULL,
                    evidence_source_ids TEXT NOT NULL DEFAULT '[]',
                    result_preview TEXT NOT NULL DEFAULT '',

                    FOREIGN KEY (interaction_id)
                        REFERENCES interactions(id)
                        ON DELETE CASCADE,

                    UNIQUE (interaction_id, sequence_number)
                );

                CREATE INDEX IF NOT EXISTS idx_tool_calls_interaction
                    ON tool_calls(interaction_id);
                """
            )
            self._migrate_interactions_schema(connection)
            self._migrate_tool_calls_schema(connection)

    @staticmethod
    def _migrate_interactions_schema(
            connection: sqlite3.Connection,
    ) -> None:
        """Add phase-6 columns to databases created by earlier versions."""

        existing_columns = {
            str(row["name"])
            for row in connection.execute(
                "PRAGMA table_info(interactions)"
            ).fetchall()
        }

        for column_name, definition in (
                INTERACTION_MIGRATION_COLUMNS.items()
        ):
            if column_name in existing_columns:
                continue

            connection.execute(
                f"ALTER TABLE interactions "
                f"ADD COLUMN {column_name} {definition}"
            )

        connection.execute(
            """
            UPDATE interactions
            SET benchmark_technical_pass = benchmark_pass
            WHERE benchmark_technical_pass IS NULL
              AND benchmark_pass IS NOT NULL
            """
        )

        release_rows = connection.execute(
            """
            SELECT
                id,
                benchmark_pass,
                benchmark_technical_pass,
                benchmark_semantic_pass,
                benchmark_release_pass,
                benchmark_failure_reasons,
                step_outputs_json,
                final_safety_policy_pass
            FROM interactions
            WHERE (
                    benchmark_release_pass IS NULL
                    AND benchmark_pass IS NOT NULL
                )
               OR benchmark_release_pass = 1
            """
        ).fetchall()

        for row in release_rows:
            safety_policy_pass = row["final_safety_policy_pass"]

            if safety_policy_pass is None:
                try:
                    step_outputs = json.loads(
                        row["step_outputs_json"] or "{}"
                    )
                except (TypeError, json.JSONDecodeError):
                    step_outputs = {}

                if isinstance(step_outputs, dict):
                    final_step = step_outputs.get(
                        "create_final_answer"
                    )
                    if isinstance(final_step, dict):
                        metrics = final_step.get("metrics")
                        if isinstance(metrics, dict):
                            additional_metrics = metrics.get(
                                "additional_metrics"
                            )
                            if isinstance(additional_metrics, dict):
                                telemetry = additional_metrics.get(
                                    WORKFLOW_TELEMETRY_KEY
                                )
                                if isinstance(telemetry, dict):
                                    final_check = telemetry.get(
                                        "final_output_check"
                                    )
                                    if (
                                        isinstance(final_check, dict)
                                        and type(
                                            final_check.get(
                                                "safety_policy_pass"
                                            )
                                        ) is bool
                                    ):
                                        safety_policy_pass = int(
                                            final_check[
                                                "safety_policy_pass"
                                            ]
                                        )
                                        connection.execute(
                                            """
                                            UPDATE interactions
                                            SET final_safety_policy_pass = ?
                                            WHERE id = ?
                                            """,
                                            (
                                                safety_policy_pass,
                                                row["id"],
                                            ),
                                        )

            legacy_without_release_evidence = (
                row["benchmark_release_pass"] is None
                and row["benchmark_pass"] is not None
            )
            invalid_existing_release = (
                row["benchmark_release_pass"] == 1
                and not (
                    row["benchmark_technical_pass"] == 1
                    and row["benchmark_semantic_pass"] == 1
                    and safety_policy_pass == 1
                )
            )

            if not (
                legacy_without_release_evidence
                or invalid_existing_release
            ):
                continue

            raw_reasons = row["benchmark_failure_reasons"]

            try:
                parsed_reasons = json.loads(raw_reasons or "[]")
            except (TypeError, json.JSONDecodeError):
                parsed_reasons = []

            reasons = (
                [
                    str(reason)
                    for reason in parsed_reasons
                    if str(reason).strip()
                ]
                if isinstance(parsed_reasons, list)
                else []
            )

            if row["benchmark_semantic_pass"] != 1:
                semantic_reason = (
                    "semantic_not_evaluated"
                    if row["benchmark_semantic_pass"] is None
                    else "semantic_check_failed"
                )
                if semantic_reason not in reasons:
                    reasons.append(semantic_reason)

            if safety_policy_pass != 1:
                safety_reason = (
                    "safety_policy_not_evaluated"
                    if safety_policy_pass is None
                    else "safety_policy_failed"
                )
                if safety_reason not in reasons:
                    reasons.append(safety_reason)

            # Legacy benchmark_pass represented the old technical-only
            # criterion. Preserve it in benchmark_technical_pass, but never
            # promote it to a release pass without semantic and safety data.
            # Existing beta rows are also repaired unless technical,
            # semantic, and persisted/embedded safety evidence are all true.
            connection.execute(
                """
                UPDATE interactions
                SET
                    benchmark_pass = 0,
                    benchmark_release_pass = 0,
                    benchmark_failure_reasons = ?
                WHERE id = ?
                """,
                (
                    json.dumps(reasons, ensure_ascii=False),
                    row["id"],
                ),
            )

    @staticmethod
    def _migrate_tool_calls_schema(
            connection: sqlite3.Connection,
    ) -> None:
        """Add reporter columns to tool telemetry from earlier databases."""

        existing_columns = {
            str(row["name"])
            for row in connection.execute(
                "PRAGMA table_info(tool_calls)"
            ).fetchall()
        }

        for column_name, definition in TOOL_CALL_MIGRATION_COLUMNS.items():
            if column_name in existing_columns:
                continue

            connection.execute(
                f"ALTER TABLE tool_calls "
                f"ADD COLUMN {column_name} {definition}"
            )

    # =========================================================
    # EXPERIMENT RUNS
    # =========================================================

    def start_run(self, config: ExperimentConfig) -> int:
        """
        Создаёт новую строку запуска и возвращает внутренний run_id.
        """

        now = datetime.now(timezone.utc).isoformat()

        with self._connect() as connection:
            next_number = connection.execute(
                """
                SELECT COALESCE(MAX(run_number), 0) + 1
                FROM experiment_runs
                """
            ).fetchone()[0]

            cursor = connection.execute(
                """
                INSERT INTO experiment_runs (
                    run_number,
                    started_at_utc,
                    model_name,
                    configuration_json,
                    notes
                )
                VALUES (?, ?, ?, ?, ?)
                """,
                (
                    next_number,
                    now,
                    config.model_name,
                    json.dumps(
                        asdict(config),
                        ensure_ascii=False,
                        sort_keys=True,
                    ),
                    config.notes,
                ),
            )

            experiment_run_id = int(cursor.lastrowid)

        self._export_after_persist()

        return experiment_run_id

    def finish_run(self, experiment_run_id: int) -> None:
        now = datetime.now(timezone.utc).isoformat()

        with self._connect() as connection:
            connection.execute(
                """
                UPDATE experiment_runs
                SET finished_at_utc = ?
                WHERE id = ?
                """,
                (now, experiment_run_id),
            )

        self._export_after_persist()

    def update_run_component_versions(
        self,
        experiment_run_id: int,
        updates: dict[str, str],
    ) -> None:
        """Persist artifact provenance discovered after a run has started."""

        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT configuration_json
                FROM experiment_runs
                WHERE id = ?
                """,
                (experiment_run_id,),
            ).fetchone()

            if row is None:
                raise ValueError(
                    f"Experiment run {experiment_run_id} does not exist"
                )

            try:
                configuration = json.loads(
                    row["configuration_json"]
                )
            except (TypeError, json.JSONDecodeError) as error:
                raise ValueError(
                    "Experiment configuration is not valid JSON."
                ) from error

            if not isinstance(configuration, dict):
                raise ValueError(
                    "Experiment configuration must be a JSON object."
                )

            component_versions = configuration.get(
                "component_versions"
            )
            if not isinstance(component_versions, dict):
                component_versions = {}

            component_versions.update({
                str(key): str(value)
                for key, value in updates.items()
            })
            configuration["component_versions"] = component_versions

            connection.execute(
                """
                UPDATE experiment_runs
                SET configuration_json = ?
                WHERE id = ?
                """,
                (
                    json.dumps(
                        configuration,
                        ensure_ascii=False,
                        sort_keys=True,
                    ),
                    experiment_run_id,
                ),
            )

        self._export_after_persist()

    def get_run_number(self, experiment_run_id: int) -> int:
        with self._connect() as connection:
            row = connection.execute(
                """
                SELECT run_number
                FROM experiment_runs
                WHERE id = ?
                """,
                (experiment_run_id,),
            ).fetchone()

        if row is None:
            raise ValueError(
                f"Experiment run {experiment_run_id} does not exist"
            )

        return int(row["run_number"])

    def _next_question_number(
        self,
        experiment_run_id: int,
    ) -> int:
        with self._connect() as connection:
            value = connection.execute(
                """
                SELECT COALESCE(MAX(question_number), 0) + 1
                FROM interactions
                WHERE experiment_run_id = ?
                """,
                (experiment_run_id,),
            ).fetchone()[0]

        return int(value)

    # =========================================================
    # RECORDING
    # =========================================================

    def record_workflow_result(
        self,
        *,
        experiment_run_id: int,
        question: str,
        workflow_response: Any,
        sources: list[dict[str, Any]] | None = None,
        error: str | None = None,
        benchmark_context: dict[str, Any] | None = None,
    ) -> bool | None:
        """
        Сохраняет завершённый WorkflowRunOutput.

        workflow_response ожидается как результат:
            workflow.run(..., stream=False)
        """

        question_number = self._next_question_number(
            experiment_run_id
        )

        answer = self._stringify(
            getattr(workflow_response, "content", "")
        )

        workflow_run_id = getattr(
            workflow_response,
            "run_id",
            None,
        )

        session_id = getattr(
            workflow_response,
            "session_id",
            None,
        )

        status = self._stringify(
            getattr(workflow_response, "status", "unknown")
        )

        metrics = getattr(
            workflow_response,
            "metrics",
            None,
        )

        metrics_dict = self._object_to_dict(metrics)

        step_outputs = self._extract_step_outputs(
            getattr(workflow_response, "step_results", None)
        )

        extracted_sources = (
            sources
            if sources is not None
            else self._find_sources_in_steps(step_outputs)
        )
        tool_calls = self._extract_tool_calls(step_outputs)
        reporting_metrics = self._extract_reporting_metrics(
            status=status,
            step_outputs=step_outputs,
            sources=extracted_sources,
            tool_calls=tool_calls,
        )
        benchmark_fields = self._build_benchmark_fields(
            benchmark_context=benchmark_context,
            reporting_metrics=reporting_metrics,
        )

        input_tokens = self._to_optional_int(
            metrics_dict.get("input_tokens")
        )
        output_tokens = self._to_optional_int(
            metrics_dict.get("output_tokens")
        )
        total_tokens = self._to_optional_int(
            metrics_dict.get("total_tokens")
        )

        duration_seconds = self._to_optional_float(
            metrics_dict.get("duration")
        )
        time_to_first_token = self._to_optional_float(
            metrics_dict.get("time_to_first_token")
        )

        with self._connect() as connection:
            cursor = connection.execute(
                """
                INSERT INTO interactions (
                    experiment_run_id,
                    question_number,
                    workflow_run_id,
                    session_id,
                    created_at_utc,
                    question,
                    answer,
                    status,
                    duration_seconds,
                    input_tokens,
                    output_tokens,
                    total_tokens,
                    time_to_first_token,
                    sources_json,
                    step_outputs_json,
                    metrics_json,
                    tool_call_count,
                    tool_names,
                    tool_error_count,
                    evidence_read_performed,
                    citation_count,
                    used_source_ids,
                    invalid_source_ids,
                    source_count,
                    researcher_success,
                    validator_success,
                    final_status,
                    final_safety_blocked,
                    final_citation_count,
                    final_used_source_ids,
                    final_invalid_source_ids,
                    final_output_policy_pass,
                    final_safety_policy_pass,
                    evidence_source_ids,
                    benchmark_scenario_id,
                    benchmark_scenario_type,
                    benchmark_turn_id,
                    benchmark_turn_number,
                    benchmark_pass,
                    benchmark_technical_pass,
                    benchmark_semantic_pass,
                    benchmark_release_pass,
                    benchmark_failure_reasons,
                    semantic_failure_reasons,
                    error
                )
                VALUES (
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?
                )
                """,
                (
                    experiment_run_id,
                    question_number,
                    workflow_run_id,
                    session_id,
                    datetime.now(timezone.utc).isoformat(),
                    question,
                    answer,
                    status,
                    duration_seconds,
                    input_tokens,
                    output_tokens,
                    total_tokens,
                    time_to_first_token,
                    json.dumps(
                        extracted_sources,
                        ensure_ascii=False,
                    ),
                    json.dumps(
                        step_outputs,
                        ensure_ascii=False,
                    ),
                    json.dumps(
                        metrics_dict,
                        ensure_ascii=False,
                        default=str,
                    ),
                    reporting_metrics["tool_call_count"],
                    json.dumps(
                        reporting_metrics["tool_names"],
                        ensure_ascii=False,
                    ),
                    reporting_metrics["tool_error_count"],
                    int(reporting_metrics["evidence_read_performed"]),
                    reporting_metrics["citation_count"],
                    json.dumps(
                        reporting_metrics["used_source_ids"],
                        ensure_ascii=False,
                    ),
                    json.dumps(
                        reporting_metrics["invalid_source_ids"],
                        ensure_ascii=False,
                    ),
                    reporting_metrics["source_count"],
                    reporting_metrics["researcher_success"],
                    reporting_metrics["validator_success"],
                    reporting_metrics["final_status"],
                    reporting_metrics["final_safety_blocked"],
                    reporting_metrics["final_citation_count"],
                    json.dumps(
                        reporting_metrics["final_used_source_ids"],
                        ensure_ascii=False,
                    ),
                    json.dumps(
                        reporting_metrics["final_invalid_source_ids"],
                        ensure_ascii=False,
                    ),
                    reporting_metrics["final_output_policy_pass"],
                    (
                        int(
                            reporting_metrics[
                                "final_safety_policy_pass"
                            ]
                        )
                        if reporting_metrics[
                            "final_safety_policy_pass"
                        ] is not None
                        else None
                    ),
                    json.dumps(
                        reporting_metrics["evidence_source_ids"],
                        ensure_ascii=False,
                    ),
                    benchmark_fields["benchmark_scenario_id"],
                    benchmark_fields["benchmark_scenario_type"],
                    benchmark_fields["benchmark_turn_id"],
                    benchmark_fields["benchmark_turn_number"],
                    benchmark_fields["benchmark_pass"],
                    benchmark_fields["benchmark_technical_pass"],
                    benchmark_fields["benchmark_semantic_pass"],
                    benchmark_fields["benchmark_release_pass"],
                    json.dumps(
                        benchmark_fields["benchmark_failure_reasons"],
                        ensure_ascii=False,
                    ),
                    json.dumps(
                        benchmark_fields["semantic_failure_reasons"],
                        ensure_ascii=False,
                    ),
                    error,
                ),
            )
            interaction_id = int(cursor.lastrowid)
            self._insert_tool_calls(
                connection,
                interaction_id=interaction_id,
                tool_calls=tool_calls,
            )

        # После каждого вопроса Excel обновляется. Ошибка производного
        # XLSX не отменяет уже подтверждённую запись в SQLite.
        self._export_after_persist()

        release_pass = benchmark_fields["benchmark_release_pass"]
        return (
            bool(release_pass)
            if release_pass is not None
            else None
        )

    def record_failure(
        self,
        *,
        experiment_run_id: int,
        question: str,
        error: Exception | str,
        session_id: str | None = None,
        benchmark_context: dict[str, Any] | None = None,
    ) -> None:
        question_number = self._next_question_number(
            experiment_run_id
        )

        error_text = str(error)
        benchmark_fields = self._build_benchmark_fields(
            benchmark_context=benchmark_context,
            reporting_metrics=None,
        )

        with self._connect() as connection:
            connection.execute(
                """
                INSERT INTO interactions (
                    experiment_run_id,
                    question_number,
                    session_id,
                    created_at_utc,
                    question,
                    answer,
                    status,
                    final_status,
                    benchmark_scenario_id,
                    benchmark_scenario_type,
                    benchmark_turn_id,
                    benchmark_turn_number,
                    benchmark_pass,
                    benchmark_technical_pass,
                    benchmark_semantic_pass,
                    benchmark_release_pass,
                    benchmark_failure_reasons,
                    semantic_failure_reasons,
                    error
                )
                VALUES (
                    ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?,
                    ?, ?
                )
                """,
                (
                    experiment_run_id,
                    question_number,
                    session_id,
                    datetime.now(timezone.utc).isoformat(),
                    question,
                    "",
                    "failed",
                    "failed",
                    benchmark_fields["benchmark_scenario_id"],
                    benchmark_fields["benchmark_scenario_type"],
                    benchmark_fields["benchmark_turn_id"],
                    benchmark_fields["benchmark_turn_number"],
                    benchmark_fields["benchmark_pass"],
                    benchmark_fields["benchmark_technical_pass"],
                    benchmark_fields["benchmark_semantic_pass"],
                    benchmark_fields["benchmark_release_pass"],
                    json.dumps(
                        benchmark_fields["benchmark_failure_reasons"],
                        ensure_ascii=False,
                    ),
                    json.dumps(
                        benchmark_fields["semantic_failure_reasons"],
                        ensure_ascii=False,
                    ),
                    error_text,
                ),
            )

        self._export_after_persist()

    # =========================================================
    # EXTRACTION
    # =========================================================

    @staticmethod
    def _object_to_dict(value: Any) -> dict[str, Any]:
        if value is None:
            return {}

        if isinstance(value, dict):
            return value

        to_dict = getattr(value, "to_dict", None)

        if callable(to_dict):
            result = to_dict()
            return result if isinstance(result, dict) else {}

        model_dump = getattr(value, "model_dump", None)

        if callable(model_dump):
            result = model_dump()
            return result if isinstance(result, dict) else {}

        if hasattr(value, "__dict__"):
            return {
                key: item
                for key, item in vars(value).items()
                if not key.startswith("_")
            }

        return {}

    def _extract_step_outputs(
        self,
        step_results: Any,
    ) -> dict[str, Any]:
        output: dict[str, Any] = {}

        if not step_results:
            return output

        def visit(items: Iterable[Any]) -> None:
            for item in items:
                if isinstance(item, list):
                    visit(item)
                    continue

                name = (
                    getattr(item, "step_name", None)
                    or getattr(item, "step_id", None)
                    or f"step_{len(output) + 1}"
                )

                output[str(name)] = {
                    "success": getattr(item, "success", None),
                    "error": getattr(item, "error", None),
                    "content": self._stringify(
                        getattr(item, "content", None)
                    ),
                    "metrics": self._object_to_dict(
                        getattr(item, "metrics", None)
                    ),
                }

        visit(step_results)

        return output

    @staticmethod
    def _find_sources_in_steps(
        step_outputs: dict[str, Any],
    ) -> list[dict[str, Any]]:
        """
        Пытается найти sources в JSON-результате шага
        collect_official_sources.
        """

        likely_names = (
            "collect_official_sources",
            "collect_sources",
            "source_collector",
        )

        for name in likely_names:
            step = step_outputs.get(name)

            if not step:
                continue

            content = step.get("content", "")

            try:
                parsed = json.loads(content)
            except (TypeError, json.JSONDecodeError):
                continue

            sources = parsed.get("sources")

            if isinstance(sources, list):
                return sources

        return []

    @staticmethod
    def _step_telemetry(step_data: dict[str, Any]) -> dict[str, Any]:
        metrics = step_data.get("metrics")

        if not isinstance(metrics, dict):
            return {}

        additional_metrics = metrics.get("additional_metrics")

        if not isinstance(additional_metrics, dict):
            return {}

        telemetry = additional_metrics.get(WORKFLOW_TELEMETRY_KEY)
        return telemetry if isinstance(telemetry, dict) else {}

    def _extract_tool_calls(
            self,
            step_outputs: dict[str, Any],
    ) -> list[dict[str, Any]]:
        """Flatten agent tool summaries while retaining step and attempt."""

        tool_calls: list[dict[str, Any]] = []

        for step_name, step_data in step_outputs.items():
            if not isinstance(step_data, dict):
                continue

            telemetry = self._step_telemetry(step_data)
            summary = telemetry.get("tool_summary")

            if not isinstance(summary, dict):
                continue

            calls = summary.get("calls")

            if not isinstance(calls, (list, tuple)):
                continue

            agent_name = str(
                telemetry.get("agent_name")
                or step_name
            )

            for call in calls:
                if not isinstance(call, dict):
                    continue

                arguments = call.get("arguments")
                if not isinstance(arguments, dict):
                    arguments = {}

                result_preview = str(
                    call.get("result_preview") or ""
                )
                evidence_read = bool(call.get("evidence_read"))
                evidence_source_ids = self._tool_evidence_source_ids(
                    call=call,
                    arguments=arguments,
                    result_preview=result_preview,
                    evidence_read=evidence_read,
                )

                tool_calls.append(
                    {
                        "sequence_number": len(tool_calls) + 1,
                        "step_name": str(step_name),
                        "agent_name": agent_name,
                        "attempt": (
                            self._to_optional_int(
                                call.get("attempt")
                            )
                            or 1
                        ),
                        "tool_name": str(
                            call.get("name") or ""
                        ),
                        "arguments": arguments,
                        "success": bool(call.get("success")),
                        "error": bool(call.get("error")),
                        "duration_seconds": self._to_optional_float(
                            call.get("duration_seconds")
                        ),
                        "evidence_read": evidence_read,
                        "evidence_source_ids": evidence_source_ids,
                        "result_preview": result_preview[:240],
                    }
                )

        return tool_calls

    @staticmethod
    def _normalize_source_ids(value: Any) -> list[str]:
        if isinstance(value, str):
            candidates: Iterable[Any] = [value]
        elif isinstance(value, (list, tuple, set)):
            candidates = value
        else:
            return []

        normalized = {
            str(candidate).strip().upper()
            for candidate in candidates
            if re.fullmatch(
                r"S\d+",
                str(candidate).strip(),
                flags=re.IGNORECASE,
            )
        }
        return sorted(
            normalized,
            key=lambda source_id: (
                int(source_id[1:]),
                source_id,
            ),
        )

    @classmethod
    def _tool_evidence_source_ids(
            cls,
            *,
            call: dict[str, Any],
            arguments: dict[str, Any],
            result_preview: str,
            evidence_read: bool,
    ) -> list[str]:
        """Extract IDs attached to an evidence-producing tool call."""

        if not evidence_read:
            return []

        source_ids: list[str] = []
        source_ids.extend(
            cls._normalize_source_ids(
                call.get("evidence_source_ids")
            )
        )
        source_ids.extend(
            cls._normalize_source_ids(arguments.get("source_id"))
        )
        source_ids.extend(
            cls._normalize_source_ids(arguments.get("source_ids"))
        )

        # Older telemetry has only a JSON result preview. Use it when it is
        # complete, while preferring the explicit evidence_source_ids field.
        try:
            parsed_preview = json.loads(result_preview)
        except (TypeError, json.JSONDecodeError):
            parsed_preview = None

        if isinstance(parsed_preview, dict):
            source_ids.extend(
                cls._normalize_source_ids(
                    parsed_preview.get("source_id")
                )
            )
            for collection_name in ("matches", "sources"):
                collection = parsed_preview.get(collection_name)
                if not isinstance(collection, list):
                    continue
                source_ids.extend(
                    cls._normalize_source_ids(
                        [
                            item.get("source_id")
                            for item in collection
                            if isinstance(item, dict)
                        ]
                    )
                )

        return cls._normalize_source_ids(source_ids)

    def _extract_reporting_metrics(
            self,
            *,
            status: str,
            step_outputs: dict[str, Any],
            sources: list[dict[str, Any]],
            tool_calls: list[dict[str, Any]],
    ) -> dict[str, Any]:
        citation_check: dict[str, Any] = {}
        citation_step = step_outputs.get(
            "check_research_citations",
            {},
        )

        if isinstance(citation_step, dict):
            citation_telemetry = self._step_telemetry(
                citation_step
            ).get("citation_check")

            if isinstance(citation_telemetry, dict):
                citation_check = citation_telemetry
            else:
                try:
                    parsed = json.loads(
                        citation_step.get("content") or "{}"
                    )
                except (TypeError, json.JSONDecodeError):
                    parsed = {}

                if isinstance(parsed, dict):
                    citation_check = parsed

        def step_success(step_name: str) -> int | None:
            step = step_outputs.get(step_name)

            if not isinstance(step, dict):
                return None

            success = step.get("success")

            if success is None:
                return None

            return int(bool(success))

        used_source_ids = citation_check.get(
            "used_source_ids",
            [],
        )
        invalid_source_ids = citation_check.get(
            "invalid_source_ids",
            [],
        )
        final_step = step_outputs.get("create_final_answer", {})
        final_safety_check: dict[str, Any] = {}
        final_output_check: dict[str, Any] = {}

        if isinstance(final_step, dict):
            final_telemetry = self._step_telemetry(final_step)
            candidate = final_telemetry.get(
                "final_safety_check"
            )
            if isinstance(candidate, dict):
                final_safety_check = candidate
            candidate = final_telemetry.get(
                "final_output_check"
            )
            if isinstance(candidate, dict):
                final_output_check = candidate

        final_citation_check = final_output_check.get(
            "citation_check"
        )
        if not isinstance(final_citation_check, dict):
            final_citation_check = final_output_check

        final_used_source_ids = self._normalize_source_ids(
            final_citation_check.get("used_source_ids")
        )
        final_invalid_source_ids = self._normalize_source_ids(
            final_citation_check.get("invalid_source_ids")
        )
        final_policy_value: Any = None
        for key in ("passed", "policy_pass"):
            if key in final_output_check:
                final_policy_value = final_output_check.get(key)
                break

        evidence_source_ids: list[str] = []
        for tool_call in tool_calls:
            evidence_source_ids.extend(
                tool_call.get("evidence_source_ids", [])
            )
        for step_data in step_outputs.values():
            if not isinstance(step_data, dict):
                continue
            telemetry = self._step_telemetry(step_data)
            evidence_source_ids.extend(
                self._normalize_source_ids(
                    telemetry.get("evidence_source_ids")
                )
            )

        return {
            "tool_call_count": len(tool_calls),
            "tool_names": [
                tool_call["tool_name"]
                for tool_call in tool_calls
            ],
            "tool_error_count": sum(
                tool_call["error"]
                for tool_call in tool_calls
            ),
            "evidence_read_performed": any(
                tool_call["evidence_read"]
                for tool_call in tool_calls
            ),
            "citation_count": (
                self._to_optional_int(
                    citation_check.get("citation_count")
                )
                or 0
            ),
            "used_source_ids": (
                list(used_source_ids)
                if isinstance(used_source_ids, (list, tuple))
                else []
            ),
            "invalid_source_ids": (
                list(invalid_source_ids)
                if isinstance(invalid_source_ids, (list, tuple))
                else []
            ),
            "source_count": len(sources),
            "researcher_success": step_success(
                "analyze_official_sources"
            ),
            "validator_success": step_success(
                "validate_research"
            ),
            "required_step_success": {
                step_name: step_success(step_name)
                for step_name in REQUIRED_WORKFLOW_STEPS
            },
            "final_status": status or "unknown",
            "final_safety_blocked": int(
                bool(final_safety_check.get("blocked"))
            ),
            "final_citation_count": (
                self._to_optional_int(
                    final_citation_check.get("citation_count")
                )
                or 0
            ),
            "final_used_source_ids": final_used_source_ids,
            "final_invalid_source_ids": final_invalid_source_ids,
            "final_output_policy_pass": (
                int(bool(final_policy_value))
                if final_policy_value is not None
                else None
            ),
            "final_output_check_present": bool(final_output_check),
            "final_citations_required": (
                bool(final_output_check.get("citations_required"))
                if "citations_required" in final_output_check
                else None
            ),
            "final_safety_policy_pass": (
                bool(final_output_check.get("safety_policy_pass"))
                if "safety_policy_pass" in final_output_check
                else None
            ),
            "evidence_source_ids": self._normalize_source_ids(
                evidence_source_ids
            ),
        }

    @staticmethod
    def _workflow_status_succeeded(status: Any) -> bool:
        normalized = str(status or "").strip().casefold()
        normalized = normalized.rsplit(".", maxsplit=1)[-1]
        return normalized in {
            "complete",
            "completed",
            "success",
            "succeeded",
        }

    def _build_benchmark_fields(
            self,
            *,
            benchmark_context: dict[str, Any] | None,
            reporting_metrics: dict[str, Any] | None,
    ) -> dict[str, Any]:
        context = (
            benchmark_context
            if isinstance(benchmark_context, dict)
            else {}
        )
        scenario_id = str(context.get("scenario_id") or "").strip()

        if not scenario_id:
            return {
                "benchmark_scenario_id": None,
                "benchmark_scenario_type": None,
                "benchmark_turn_id": None,
                "benchmark_turn_number": None,
                "benchmark_pass": None,
                "benchmark_technical_pass": None,
                "benchmark_semantic_pass": None,
                "benchmark_release_pass": None,
                "benchmark_failure_reasons": [],
                "semantic_failure_reasons": [],
            }

        technical_failure_reasons: list[str] = []

        if reporting_metrics is None:
            technical_failure_reasons.append("workflow_failed")
        else:
            if not self._workflow_status_succeeded(
                    reporting_metrics.get("final_status")
            ):
                technical_failure_reasons.append(
                    "workflow_not_successful"
                )

            if int(reporting_metrics.get("source_count") or 0) <= 0:
                technical_failure_reasons.append("no_sources")

            if not bool(
                    reporting_metrics.get("evidence_read_performed")
            ):
                technical_failure_reasons.append(
                    "evidence_not_read"
                )

            required_step_success = reporting_metrics.get(
                "required_step_success"
            )
            if not isinstance(required_step_success, dict):
                required_step_success = {}

            step_failure_reasons = {
                "build_search_query": "query_builder_not_successful",
                "collect_official_sources": (
                    "source_collection_not_successful"
                ),
                "analyze_official_sources": (
                    "researcher_not_successful"
                ),
                "check_research_citations": (
                    "citation_check_not_successful"
                ),
                "validate_research": "validator_not_successful",
                "create_final_answer": (
                    "final_answer_not_successful"
                ),
            }

            for step_name in REQUIRED_WORKFLOW_STEPS:
                if required_step_success.get(step_name) != 1:
                    technical_failure_reasons.append(
                        step_failure_reasons[step_name]
                    )

            if reporting_metrics.get("final_output_check_present"):
                if reporting_metrics.get(
                        "final_output_policy_pass"
                ) != 1:
                    technical_failure_reasons.append(
                        "final_output_policy_failed"
                    )

                if reporting_metrics.get(
                        "final_invalid_source_ids"
                ):
                    technical_failure_reasons.append(
                        "final_invalid_source_ids"
                    )

                citations_required = reporting_metrics.get(
                    "final_citations_required"
                )
                if citations_required is None:
                    citations_required = self._optional_bool_value(
                        context.get("citations_required")
                    )
                if citations_required is None:
                    citations_required = True

                if (
                    citations_required
                    and int(
                        reporting_metrics.get(
                            "final_citation_count"
                        )
                        or 0
                    ) <= 0
                ):
                    technical_failure_reasons.append(
                        "no_final_citations"
                    )
            else:
                # Compatibility for reports produced before final-output
                # telemetry existed.
                if reporting_metrics.get("invalid_source_ids"):
                    technical_failure_reasons.append(
                        "invalid_source_ids"
                    )

                if int(
                    reporting_metrics.get("citation_count") or 0
                ) <= 0:
                    technical_failure_reasons.append("no_citations")

        semantic_pass = self._optional_bool_value(
            context.get("semantic_pass")
        )
        raw_semantic_reasons = context.get(
            "semantic_failure_reasons",
            [],
        )
        semantic_failure_reasons = (
            [
                str(reason)
                for reason in raw_semantic_reasons
                if str(reason).strip()
            ]
            if isinstance(raw_semantic_reasons, (list, tuple))
            else []
        )
        if semantic_pass is False and not semantic_failure_reasons:
            semantic_failure_reasons.append(
                "semantic_expectations_failed"
            )

        safety_policy_pass = self._optional_bool_value(
            context.get("safety_policy_pass")
        )
        if (
            safety_policy_pass is None
            and reporting_metrics is not None
        ):
            safety_policy_pass = self._optional_bool_value(
                reporting_metrics.get("final_safety_policy_pass")
            )

        technical_pass = not technical_failure_reasons
        release_pass = (
            technical_pass
            and semantic_pass is True
            and safety_policy_pass is True
        )
        release_failure_reasons = list(
            technical_failure_reasons
        )
        if semantic_pass is False:
            release_failure_reasons.append(
                "semantic_check_failed"
            )
        elif semantic_pass is None:
            release_failure_reasons.append(
                "semantic_not_evaluated"
            )
        if safety_policy_pass is False:
            release_failure_reasons.append(
                "safety_policy_failed"
            )
        elif safety_policy_pass is None:
            release_failure_reasons.append(
                "safety_policy_not_evaluated"
            )

        return {
            "benchmark_scenario_id": scenario_id,
            "benchmark_scenario_type": (
                str(context.get("scenario_type") or "").strip()
                or None
            ),
            "benchmark_turn_id": (
                str(context.get("turn_id") or "").strip()
                or None
            ),
            "benchmark_turn_number": self._to_optional_int(
                context.get("turn_number")
            ),
            # benchmark_pass remains the compatibility alias for release pass.
            "benchmark_pass": int(release_pass),
            "benchmark_technical_pass": int(technical_pass),
            "benchmark_semantic_pass": (
                int(semantic_pass)
                if semantic_pass is not None
                else None
            ),
            "benchmark_release_pass": int(release_pass),
            "benchmark_failure_reasons": release_failure_reasons,
            "semantic_failure_reasons": semantic_failure_reasons,
        }

    @staticmethod
    def _insert_tool_calls(
            connection: sqlite3.Connection,
            *,
            interaction_id: int,
            tool_calls: list[dict[str, Any]],
    ) -> None:
        for tool_call in tool_calls:
            connection.execute(
                """
                INSERT INTO tool_calls (
                    interaction_id,
                    sequence_number,
                    step_name,
                    agent_name,
                    attempt,
                    tool_name,
                    arguments_json,
                    success,
                    error,
                    duration_seconds,
                    evidence_read,
                    evidence_source_ids,
                    result_preview
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    interaction_id,
                    tool_call["sequence_number"],
                    tool_call["step_name"],
                    tool_call["agent_name"],
                    tool_call["attempt"],
                    tool_call["tool_name"],
                    json.dumps(
                        tool_call["arguments"],
                        ensure_ascii=False,
                        default=str,
                    ),
                    int(tool_call["success"]),
                    int(tool_call["error"]),
                    tool_call["duration_seconds"],
                    int(tool_call["evidence_read"]),
                    json.dumps(
                        tool_call["evidence_source_ids"],
                        ensure_ascii=False,
                    ),
                    tool_call["result_preview"],
                ),
            )

    @staticmethod
    def _stringify(value: Any) -> str:
        if value is None:
            return ""

        if isinstance(value, str):
            return value

        if isinstance(value, (dict, list)):
            return json.dumps(
                value,
                ensure_ascii=False,
                indent=2,
            )

        return str(value)

    @staticmethod
    def _to_optional_int(value: Any) -> int | None:
        try:
            return int(value) if value is not None else None
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _to_optional_float(value: Any) -> float | None:
        try:
            return float(value) if value is not None else None
        except (TypeError, ValueError):
            return None

    @staticmethod
    def _optional_bool_value(value: Any) -> bool | None:
        if value is None:
            return None
        if isinstance(value, bool):
            return value
        if isinstance(value, (int, float)):
            return bool(value)
        if isinstance(value, str):
            normalized = value.strip().casefold()
            if normalized in {"true", "yes", "1", "pass", "passed"}:
                return True
            if normalized in {"false", "no", "0", "fail", "failed"}:
                return False
        return None

    # =========================================================
    # XLSX EXPORT
    # =========================================================

    def _export_after_persist(self) -> None:
        """
        Refresh the derived workbook without obscuring a committed DB write.

        Callers can inspect ``last_export_error`` and retry ``export_xlsx``.
        A warning is deliberately used instead of an exception: otherwise a
        runner may incorrectly record a second, failed interaction after the
        successful interaction has already committed.
        """

        try:
            self.export_xlsx()
        except Exception as error:
            self.last_export_error = (
                f"{type(error).__name__}: {error}"
            )
            warnings.warn(
                (
                    "SQLite data was persisted, but the derived XLSX "
                    f"export failed: {self.last_export_error}"
                ),
                ReportExportWarning,
                stacklevel=2,
            )
        else:
            self.last_export_error = None

    def export_xlsx(self) -> None:
        """
        Не изменяет Excel вручную строка за строкой.

        Каждый раз читает полные данные из SQLite и строит XLSX заново.
        Поэтому SQLite остаётся единственным источником истины.
        """

        runs, interactions, tool_calls = self._load_report_data()

        workbook = Workbook()

        summary_sheet = workbook.active
        summary_sheet.title = "Benchmark Summary"

        runs_sheet = workbook.create_sheet("Runs")
        details_sheet = workbook.create_sheet("Interactions")
        steps_sheet = workbook.create_sheet("Step Outputs")
        tool_calls_sheet = workbook.create_sheet("Tool Calls")

        self._write_benchmark_summary_sheet(
            summary_sheet,
            runs,
            interactions,
        )
        self._write_runs_sheet(
            runs_sheet,
            runs,
            interactions,
        )

        self._write_interactions_sheet(
            details_sheet,
            runs,
            interactions,
        )

        self._write_steps_sheet(
            steps_sheet,
            runs,
            interactions,
        )
        self._write_tool_calls_sheet(
            tool_calls_sheet,
            runs,
            interactions,
            tool_calls,
        )

        self._save_atomically(workbook)
        self.last_export_error = None

    def _load_report_data(
        self,
    ) -> tuple[
        list[sqlite3.Row],
        list[sqlite3.Row],
        list[sqlite3.Row],
    ]:
        with self._connect() as connection:
            runs = connection.execute(
                """
                SELECT *
                FROM experiment_runs
                ORDER BY run_number
                """
            ).fetchall()

            interactions = connection.execute(
                """
                SELECT *
                FROM interactions
                ORDER BY experiment_run_id, question_number
                """
            ).fetchall()

            tool_calls = connection.execute(
                """
                SELECT *
                FROM tool_calls
                ORDER BY interaction_id, sequence_number
                """
            ).fetchall()

        return runs, interactions, tool_calls

    def _write_benchmark_summary_sheet(
        self,
        sheet: Any,
        runs: list[sqlite3.Row],
        interactions: list[sqlite3.Row],
    ) -> None:
        """Build a compact, presentation-ready model comparison."""

        by_run: dict[int, list[sqlite3.Row]] = {}
        for interaction in interactions:
            by_run.setdefault(
                interaction["experiment_run_id"],
                [],
            ).append(interaction)

        headers = [
            "Run",
            "Model",
            "Benchmark Version",
            "Config Digest",
            "Git Commit",
            "Questions",
            "Technical Passes",
            "Technical Failures",
            "Technical Pass Rate",
            "Semantic Evaluated",
            "Semantic Passes",
            "Semantic Failures",
            "Semantic Pass Rate",
            "Release Passes",
            "Release Failures",
            "Release Pass Rate",
            "Workflow Failures",
            "Tool Calls",
            "Tool Errors",
            "Evidence Reads",
            "Evidence Source IDs",
            "Researcher Citations",
            "Final Citations",
            "Researcher Invalid Source IDs",
            "Final Invalid Source IDs",
            "Final Policy Passes",
            "Safety Policy Passes",
            "Safety Blocks",
            "Avg Duration",
        ]
        self._append_xlsx_row(sheet, headers)

        for run in runs:
            try:
                configuration = json.loads(
                    run["configuration_json"] or "{}"
                )
            except json.JSONDecodeError:
                configuration = {}

            if not isinstance(configuration, dict):
                configuration = {}

            run_interactions = by_run.get(run["id"], [])
            benchmark_interactions = [
                interaction
                for interaction in run_interactions
                if interaction["benchmark_scenario_id"] is not None
            ]
            measured = (
                benchmark_interactions
                if benchmark_interactions
                else run_interactions
            )
            technical_pass_count = sum(
                interaction["benchmark_technical_pass"] == 1
                for interaction in benchmark_interactions
            )
            technical_failure_count = sum(
                interaction["benchmark_technical_pass"] == 0
                for interaction in benchmark_interactions
            )
            semantic_interactions = [
                interaction
                for interaction in benchmark_interactions
                if interaction["benchmark_semantic_pass"] is not None
            ]
            semantic_pass_count = sum(
                interaction["benchmark_semantic_pass"] == 1
                for interaction in semantic_interactions
            )
            semantic_failure_count = sum(
                interaction["benchmark_semantic_pass"] == 0
                for interaction in semantic_interactions
            )
            release_pass_count = sum(
                interaction["benchmark_release_pass"] == 1
                for interaction in benchmark_interactions
            )
            release_failure_count = sum(
                interaction["benchmark_release_pass"] == 0
                for interaction in benchmark_interactions
            )
            durations = [
                float(interaction["duration_seconds"])
                for interaction in measured
                if interaction["duration_seconds"] is not None
            ]
            evidence_source_ids = self._normalize_source_ids([
                source_id
                for interaction in measured
                for source_id in self._json_list_value(
                    interaction["evidence_source_ids"]
                )
            ])

            self._append_xlsx_row(
                sheet,
                [
                    run["run_number"],
                    run["model_name"],
                    configuration.get("benchmark_version", ""),
                    configuration.get("benchmark_config_digest", ""),
                    configuration.get("git_commit", ""),
                    len(measured),
                    technical_pass_count,
                    technical_failure_count,
                    (
                        technical_pass_count
                        / len(benchmark_interactions)
                        if benchmark_interactions
                        else None
                    ),
                    len(semantic_interactions),
                    semantic_pass_count,
                    semantic_failure_count,
                    (
                        semantic_pass_count
                        / len(semantic_interactions)
                        if semantic_interactions
                        else None
                    ),
                    release_pass_count,
                    release_failure_count,
                    (
                        release_pass_count
                        / len(benchmark_interactions)
                        if benchmark_interactions
                        else None
                    ),
                    sum(
                        not self._workflow_status_succeeded(
                            interaction["final_status"]
                        )
                        for interaction in measured
                    ),
                    sum(
                        int(interaction["tool_call_count"] or 0)
                        for interaction in measured
                    ),
                    sum(
                        int(interaction["tool_error_count"] or 0)
                        for interaction in measured
                    ),
                    sum(
                        bool(interaction["evidence_read_performed"])
                        for interaction in measured
                    ),
                    ", ".join(evidence_source_ids),
                    sum(
                        int(interaction["citation_count"] or 0)
                        for interaction in measured
                    ),
                    sum(
                        int(interaction["final_citation_count"] or 0)
                        for interaction in measured
                    ),
                    sum(
                        len(self._json_list_value(
                            interaction["invalid_source_ids"]
                        ))
                        for interaction in measured
                    ),
                    sum(
                        len(self._json_list_value(
                            interaction["final_invalid_source_ids"]
                        ))
                        for interaction in measured
                    ),
                    sum(
                        interaction["final_output_policy_pass"] == 1
                        for interaction in measured
                    ),
                    sum(
                        interaction["final_safety_policy_pass"] == 1
                        for interaction in measured
                    ),
                    sum(
                        bool(interaction["final_safety_blocked"])
                        for interaction in measured
                    ),
                    (
                        sum(durations) / len(durations)
                        if durations
                        else None
                    ),
                ],
            )

        self._format_sheet(
            sheet,
            freeze_cell="A2",
            answer_columns=False,
        )

        for header in (
            "Technical Pass Rate",
            "Semantic Pass Rate",
            "Release Pass Rate",
        ):
            pass_rate_column = get_column_letter(
                headers.index(header) + 1
            )
            for cell in sheet[pass_rate_column][1:]:
                cell.number_format = "0.0%"

    def _write_runs_sheet(
        self,
        sheet: Any,
        runs: list[sqlite3.Row],
        interactions: list[sqlite3.Row],
    ) -> None:
        by_run: dict[int, list[sqlite3.Row]] = {}

        for interaction in interactions:
            by_run.setdefault(
                interaction["experiment_run_id"],
                [],
            ).append(interaction)

        max_questions = max(
            (
                len(items)
                for items in by_run.values()
            ),
            default=0,
        )

        headers = [
            "Run",
            "Started UTC",
            "Finished UTC",
            "Model",
            "Configuration",
            "Notes",
        ]

        for number in range(1, max_questions + 1):
            headers.extend(
                [
                    f"Q{number}",
                    f"A{number}",
                    f"Duration {number}",
                    f"Tokens {number}",
                    f"Sources {number}",
                    f"Status {number}",
                ]
            )

        self._append_xlsx_row(sheet, headers)

        for run in runs:
            row: list[Any] = [
                run["run_number"],
                run["started_at_utc"],
                run["finished_at_utc"],
                run["model_name"],
                run["configuration_json"],
                run["notes"],
            ]

            run_interactions = by_run.get(run["id"], [])

            for interaction in run_interactions:
                row.extend(
                    [
                        interaction["question"],
                        interaction["answer"],
                        interaction["duration_seconds"],
                        interaction["total_tokens"],
                        self._source_urls(
                            interaction["sources_json"]
                        ),
                        interaction["status"],
                    ]
                )

            missing = max_questions - len(run_interactions)

            for _ in range(missing):
                row.extend(["", "", "", "", "", ""])

            self._append_xlsx_row(sheet, row)

        self._format_sheet(
            sheet,
            freeze_cell="A2",
            answer_columns=True,
        )

    def _write_interactions_sheet(
        self,
        sheet: Any,
        runs: list[sqlite3.Row],
        interactions: list[sqlite3.Row],
    ) -> None:
        run_by_id = {
            run["id"]: run
            for run in runs
        }

        headers = [
            "Run",
            "Question #",
            "Model",
            "Workflow Run ID",
            "Session ID",
            "Created UTC",
            "Question",
            "Answer",
            "Status",
            "Duration",
            "Input Tokens",
            "Output Tokens",
            "Total Tokens",
            "Time To First Token",
            "Sources",
            "Tool Call Count",
            "Tool Names",
            "Tool Error Count",
            "Evidence Read",
            "Evidence Source IDs",
            "Researcher Citation Count",
            "Used Source IDs",
            "Invalid Source IDs",
            "Final Citation Count",
            "Final Used Source IDs",
            "Final Invalid Source IDs",
            "Final Output Policy Pass",
            "Final Safety Policy Pass",
            "Source Count",
            "Researcher Success",
            "Validator Success",
            "Final Status",
            "Safety Blocked",
            "Benchmark Scenario",
            "Scenario Type",
            "Turn ID",
            "Turn #",
            "Technical Pass",
            "Semantic Pass",
            "Release Pass",
            "Benchmark Failure Reasons",
            "Semantic Failure Reasons",
            "Error",
        ]

        self._append_xlsx_row(sheet, headers)

        for interaction in interactions:
            run = run_by_id[interaction["experiment_run_id"]]

            self._append_xlsx_row(
                sheet,
                [
                    run["run_number"],
                    interaction["question_number"],
                    run["model_name"],
                    interaction["workflow_run_id"],
                    interaction["session_id"],
                    interaction["created_at_utc"],
                    interaction["question"],
                    interaction["answer"],
                    interaction["status"],
                    interaction["duration_seconds"],
                    interaction["input_tokens"],
                    interaction["output_tokens"],
                    interaction["total_tokens"],
                    interaction["time_to_first_token"],
                    self._source_urls(
                        interaction["sources_json"]
                    ),
                    interaction["tool_call_count"],
                    self._json_list_text(
                        interaction["tool_names"]
                    ),
                    interaction["tool_error_count"],
                    bool(interaction["evidence_read_performed"]),
                    self._json_list_text(
                        interaction["evidence_source_ids"]
                    ),
                    interaction["citation_count"],
                    self._json_list_text(
                        interaction["used_source_ids"]
                    ),
                    self._json_list_text(
                        interaction["invalid_source_ids"]
                    ),
                    interaction["final_citation_count"],
                    self._json_list_text(
                        interaction["final_used_source_ids"]
                    ),
                    self._json_list_text(
                        interaction["final_invalid_source_ids"]
                    ),
                    self._optional_bool(
                        interaction["final_output_policy_pass"]
                    ),
                    self._optional_bool(
                        interaction["final_safety_policy_pass"]
                    ),
                    interaction["source_count"],
                    self._optional_bool(
                        interaction["researcher_success"]
                    ),
                    self._optional_bool(
                        interaction["validator_success"]
                    ),
                    interaction["final_status"],
                    bool(interaction["final_safety_blocked"]),
                    interaction["benchmark_scenario_id"],
                    interaction["benchmark_scenario_type"],
                    interaction["benchmark_turn_id"],
                    interaction["benchmark_turn_number"],
                    self._optional_bool(
                        interaction["benchmark_technical_pass"]
                    ),
                    self._optional_bool(
                        interaction["benchmark_semantic_pass"]
                    ),
                    self._optional_bool(
                        interaction["benchmark_release_pass"]
                    ),
                    self._json_list_text(
                        interaction["benchmark_failure_reasons"]
                    ),
                    self._json_list_text(
                        interaction["semantic_failure_reasons"]
                    ),
                    interaction["error"],
                ],
            )

        self._format_sheet(
            sheet,
            freeze_cell="A2",
            answer_columns=True,
        )

    def _write_tool_calls_sheet(
        self,
        sheet: Any,
        runs: list[sqlite3.Row],
        interactions: list[sqlite3.Row],
        tool_calls: list[sqlite3.Row],
    ) -> None:
        run_by_id = {
            run["id"]: run
            for run in runs
        }
        interaction_by_id = {
            interaction["id"]: interaction
            for interaction in interactions
        }

        self._append_xlsx_row(
            sheet,
            [
                "Run",
                "Question #",
                "Model",
                "Sequence",
                "Step",
                "Agent",
                "Attempt",
                "Tool",
                "Arguments",
                "Success",
                "Error",
                "Duration",
                "Evidence Read",
                "Evidence Source IDs",
                "Result Preview",
            ],
        )

        for tool_call in tool_calls:
            interaction = interaction_by_id.get(
                tool_call["interaction_id"]
            )

            if interaction is None:
                continue

            run = run_by_id[interaction["experiment_run_id"]]

            self._append_xlsx_row(
                sheet,
                [
                    run["run_number"],
                    interaction["question_number"],
                    run["model_name"],
                    tool_call["sequence_number"],
                    tool_call["step_name"],
                    tool_call["agent_name"],
                    tool_call["attempt"],
                    tool_call["tool_name"],
                    tool_call["arguments_json"],
                    bool(tool_call["success"]),
                    bool(tool_call["error"]),
                    tool_call["duration_seconds"],
                    bool(tool_call["evidence_read"]),
                    self._json_list_text(
                        tool_call["evidence_source_ids"]
                    ),
                    tool_call["result_preview"],
                ],
            )

        self._format_sheet(
            sheet,
            freeze_cell="A2",
            answer_columns=True,
        )

    def _write_steps_sheet(
        self,
        sheet: Any,
        runs: list[sqlite3.Row],
        interactions: list[sqlite3.Row],
    ) -> None:
        run_by_id = {
            run["id"]: run
            for run in runs
        }

        self._append_xlsx_row(
            sheet,
            [
                "Run",
                "Question #",
                "Model",
                "Step",
                "Success",
                "Error",
                "Content",
                "Metrics",
            ],
        )

        for interaction in interactions:
            run = run_by_id[interaction["experiment_run_id"]]

            try:
                steps = json.loads(
                    interaction["step_outputs_json"] or "{}"
                )
            except json.JSONDecodeError:
                steps = {}

            for step_name, step_data in steps.items():
                self._append_xlsx_row(
                    sheet,
                    [
                        run["run_number"],
                        interaction["question_number"],
                        run["model_name"],
                        step_name,
                        step_data.get("success"),
                        step_data.get("error"),
                        step_data.get("content"),
                        json.dumps(
                            self._step_metrics_for_xlsx(
                                step_data
                            ),
                            ensure_ascii=False,
                        ),
                    ],
                )

        self._format_sheet(
            sheet,
            freeze_cell="A2",
            answer_columns=True,
        )

    def _step_metrics_for_xlsx(
            self,
            step_data: dict[str, Any],
    ) -> dict[str, Any]:
        """Keep Step Outputs compact; tool details have their own sheet."""

        metrics = step_data.get("metrics")

        if not isinstance(metrics, dict):
            return {}

        telemetry = self._step_telemetry(step_data)
        summary = telemetry.get("tool_summary")

        if not isinstance(summary, dict):
            return metrics

        calls = summary.get("calls")
        call_count = (
            len(calls)
            if isinstance(calls, (list, tuple))
            else 0
        )
        compact_summary = {
            key: value
            for key, value in summary.items()
            if key != "calls"
        }
        compact_summary.setdefault("total_count", call_count)

        compact_telemetry = {
            key: value
            for key, value in telemetry.items()
            if key != "tool_summary"
        }
        compact_telemetry["tool_summary"] = compact_summary

        compact_metrics = {
            key: value
            for key, value in metrics.items()
            if key != "additional_metrics"
        }
        compact_metrics["additional_metrics"] = {
            WORKFLOW_TELEMETRY_KEY: compact_telemetry,
        }

        return compact_metrics

    @staticmethod
    def _source_urls(sources_json: str) -> str:
        try:
            sources = json.loads(sources_json or "[]")
        except json.JSONDecodeError:
            return ""

        urls: list[str] = []

        for source in sources:
            if not isinstance(source, dict):
                continue

            source_id = source.get("source_id", "")
            url = source.get("url", "")

            if url:
                urls.append(
                    f"{source_id}: {url}"
                    if source_id
                    else url
                )

        return "\n".join(urls)

    @staticmethod
    def _json_list_value(value: str) -> list[Any]:
        try:
            parsed = json.loads(value or "[]")
        except (TypeError, json.JSONDecodeError):
            return []

        if not isinstance(parsed, list):
            return []

        return parsed

    @classmethod
    def _json_list_text(cls, value: str) -> str:
        return ", ".join(
            str(item)
            for item in cls._json_list_value(value)
        )

    @staticmethod
    def _optional_bool(value: Any) -> bool | None:
        if value is None:
            return None

        return bool(value)

    @staticmethod
    def _sanitize_xlsx_value(value: Any) -> Any:
        """
        Keep untrusted strings literal and valid in an XLSX XML cell.

        Spreadsheet programs may interpret =, +, -, or @ as formulas even
        when an attacker hides the prefix behind whitespace/control
        characters. A leading apostrophe is the conventional literal marker.
        Numeric and boolean values retain their native cell types.
        """

        if not isinstance(value, str):
            return value

        first_meaningful = next(
            (
                character
                for character in value
                if not (
                    character.isspace()
                    or unicodedata.category(character) in {"Cc", "Cf"}
                )
            ),
            "",
        )
        sanitized = ILLEGAL_XLSX_CONTROL_RE.sub("\uFFFD", value)

        if first_meaningful in {"=", "+", "-", "@"}:
            return "'" + sanitized

        return sanitized

    @classmethod
    def _append_xlsx_row(
            cls,
            sheet: Any,
            values: Iterable[Any],
    ) -> None:
        sheet.append([
            cls._sanitize_xlsx_value(value)
            for value in values
        ])

    @staticmethod
    def _format_sheet(
        sheet: Any,
        *,
        freeze_cell: str,
        answer_columns: bool,
    ) -> None:
        sheet.freeze_panes = freeze_cell
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 30

        for cell in sheet[1]:
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78",
            )
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        for index, cell in enumerate(sheet[1], start=1):
            header = str(cell.value or "")
            column_letter = get_column_letter(index)

            if (
                    len(header) > 1
                    and header[0] in {"Q", "A"}
                    and header[1:].isdigit()
            ):
                width = 55
            elif header in {
                "Question",
                "Answer",
                "Content",
                "Configuration",
                "Metrics",
                "Sources",
                "Arguments",
                "Result Preview",
            }:
                width = 70
            elif header == "Tool Names":
                width = 45
            elif header in {
                "Used Source IDs",
                "Invalid Source IDs",
                "Final Used Source IDs",
                "Final Invalid Source IDs",
                "Evidence Source IDs",
                "Benchmark Failure Reasons",
                "Semantic Failure Reasons",
            }:
                width = 24
            elif header.startswith("Sources "):
                width = 45
            else:
                width = min(max(len(header) + 3, 12), 28)

            sheet.column_dimensions[
                column_letter
            ].width = width

            if (
                    "Duration" in header
                    or header == "Time To First Token"
            ):
                for data_cell in sheet[column_letter][1:]:
                    data_cell.number_format = "0.000"
            elif (
                    header.endswith("Count")
                    or header in {
                        "Run",
                        "Question #",
                        "Turn #",
                        "Sequence",
                        "Attempt",
                        "Input Tokens",
                        "Output Tokens",
                        "Total Tokens",
                        "Source Count",
                    }
            ):
                for data_cell in sheet[column_letter][1:]:
                    data_cell.number_format = "#,##0"

    def _save_atomically(self, workbook: Workbook) -> None:
        """
        Сначала сохраняет временный XLSX, затем заменяет отчёт.
        Старый отчёт не уничтожается при ошибке записи.
        """

        target_directory = self.xlsx_path.parent

        with tempfile.NamedTemporaryFile(
            prefix="visa_report_",
            suffix=".xlsx",
            dir=target_directory,
            delete=False,
        ) as temporary:
            temporary_path = Path(temporary.name)

        try:
            workbook.save(temporary_path)
            os.replace(
                temporary_path,
                self.xlsx_path,
            )
        finally:
            if temporary_path.exists():
                temporary_path.unlink()
