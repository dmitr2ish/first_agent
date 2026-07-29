from __future__ import annotations

import json
import sqlite3
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace

from agno.metrics import RunMetrics
from agno.workflow.step import StepOutput
from openpyxl import load_workbook

from experiment_reporter import ExperimentConfig, ExperimentReporter


def telemetry_metrics(payload: dict) -> RunMetrics:
    return RunMetrics(
        additional_metrics={
            "workflow_telemetry": payload,
        }
    )


def tool_call(
        name: str,
        *,
        attempt: int,
        success: bool,
        error: bool,
        evidence_read: bool,
        result_preview: str,
) -> dict:
    return {
        "name": name,
        "arguments": {"source_id": "S1"},
        "attempt": attempt,
        "success": success,
        "error": error,
        "duration_seconds": 0.25,
        "evidence_read": evidence_read,
        "result_preview": result_preview,
    }


class ExperimentReporterTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_directory = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_directory.name)
        self.db_path = self.root / "report.db"
        self.xlsx_path = self.root / "report.xlsx"

    def tearDown(self) -> None:
        self.temp_directory.cleanup()

    @staticmethod
    def config() -> ExperimentConfig:
        return ExperimentConfig(
            model_name="test-model",
            model_options={"temperature": 0},
            workflow_version="1.0.0",
            prompt_version="test-prompts",
            source_collector_version="test-sources",
            max_source_pages=5,
            allowed_domains=["uscis.gov"],
            git_commit="abc1234",
            notes="phase 6 test",
        )

    @staticmethod
    def workflow_response() -> SimpleNamespace:
        researcher_calls = [
            tool_call(
                "search_official_sources",
                attempt=1,
                success=False,
                error=True,
                evidence_read=False,
                result_preview="search failed",
            ),
            tool_call(
                "read_official_source",
                attempt=2,
                success=True,
                error=False,
                evidence_read=True,
                result_preview="x" * 700,
            ),
        ]
        validator_calls = [
            tool_call(
                "read_official_source",
                attempt=1,
                success=True,
                error=False,
                evidence_read=True,
                result_preview="validated source",
            ),
        ]
        source_package = {
            "sources": [
                {
                    "source_id": "S1",
                    "url": "https://uscis.gov/example-1",
                    "content": "first source",
                },
                {
                    "source_id": "S2",
                    "url": "https://travel.state.gov/example-2",
                    "content": "second source",
                },
            ],
        }
        citation_check = {
            "citation_count": 2,
            "used_source_ids": ["S1", "S99"],
            "invalid_source_ids": ["S99"],
            "allowed_source_ids": ["S1", "S2"],
            "has_citations": True,
        }

        return SimpleNamespace(
            content="final answer",
            run_id="workflow-run-1",
            session_id="session-1",
            status="completed",
            metrics={
                "input_tokens": 10,
                "output_tokens": 20,
                "total_tokens": 30,
                "duration": 1.5,
            },
            step_results=[
                StepOutput(
                    step_name="collect_official_sources",
                    content=json.dumps(source_package),
                    success=True,
                ),
                StepOutput(
                    step_name="analyze_official_sources",
                    content="research [S1]",
                    success=True,
                    metrics=telemetry_metrics(
                        {
                            "agent_name": "researcher",
                            "attempt_count": 2,
                            "tool_summary": {
                                "calls": researcher_calls,
                            },
                        }
                    ),
                ),
                StepOutput(
                    step_name="check_research_citations",
                    content=json.dumps(citation_check),
                    success=True,
                    metrics=telemetry_metrics(
                        {
                            "citation_check": {
                                **citation_check,
                                "used_source_ids": ("S1", "S99"),
                                "invalid_source_ids": ("S99",),
                            },
                        }
                    ),
                ),
                StepOutput(
                    step_name="validate_research",
                    content="validated [S1]",
                    success=True,
                    metrics=telemetry_metrics(
                        {
                            "agent_name": "validator",
                            "attempt_count": 1,
                            "tool_summary": {
                                "calls": validator_calls,
                            },
                        }
                    ),
                ),
                StepOutput(
                    step_name="create_final_answer",
                    content="final answer",
                    success=True,
                ),
            ],
        )

    def populated_reporter(self) -> tuple[ExperimentReporter, int]:
        reporter = ExperimentReporter(
            db_path=str(self.db_path),
            xlsx_path=str(self.xlsx_path),
        )
        experiment_run_id = reporter.start_run(self.config())
        reporter.record_workflow_result(
            experiment_run_id=experiment_run_id,
            question="Test question?",
            workflow_response=self.workflow_response(),
        )
        return reporter, experiment_run_id

    def test_records_interaction_tool_and_citation_metrics(self) -> None:
        self.populated_reporter()

        with sqlite3.connect(self.db_path) as connection:
            connection.row_factory = sqlite3.Row
            interaction = connection.execute(
                "SELECT * FROM interactions"
            ).fetchone()

        self.assertEqual(interaction["tool_call_count"], 3)
        self.assertEqual(
            json.loads(interaction["tool_names"]),
            [
                "search_official_sources",
                "read_official_source",
                "read_official_source",
            ],
        )
        self.assertEqual(interaction["tool_error_count"], 1)
        self.assertEqual(interaction["evidence_read_performed"], 1)
        self.assertEqual(interaction["citation_count"], 2)
        self.assertEqual(
            json.loads(interaction["used_source_ids"]),
            ["S1", "S99"],
        )
        self.assertEqual(
            json.loads(interaction["invalid_source_ids"]),
            ["S99"],
        )
        self.assertEqual(interaction["source_count"], 2)
        self.assertEqual(interaction["researcher_success"], 1)
        self.assertEqual(interaction["validator_success"], 1)
        self.assertEqual(interaction["final_status"], "completed")

    def test_records_multiple_tool_calls_and_tool_error(self) -> None:
        self.populated_reporter()

        with sqlite3.connect(self.db_path) as connection:
            connection.row_factory = sqlite3.Row
            calls = connection.execute(
                """
                SELECT *
                FROM tool_calls
                ORDER BY sequence_number
                """
            ).fetchall()

        self.assertEqual(len(calls), 3)
        self.assertEqual(
            [call["agent_name"] for call in calls],
            ["researcher", "researcher", "validator"],
        )
        self.assertEqual([call["attempt"] for call in calls], [1, 2, 1])
        self.assertEqual(calls[0]["error"], 1)
        self.assertEqual(calls[0]["success"], 0)
        self.assertEqual(calls[1]["evidence_read"], 1)
        self.assertLessEqual(len(calls[1]["result_preview"]), 240)

    def test_xlsx_is_rebuilt_from_sqlite_with_all_sheets(self) -> None:
        reporter, _ = self.populated_reporter()
        self.xlsx_path.unlink()

        reporter.export_xlsx()

        workbook = load_workbook(
            self.xlsx_path,
            read_only=False,
        )
        self.assertEqual(
            workbook.sheetnames,
            [
                "Runs",
                "Interactions",
                "Step Outputs",
                "Tool Calls",
            ],
        )
        self.assertEqual(workbook["Runs"].max_row, 2)
        self.assertEqual(workbook["Interactions"].max_row, 2)
        self.assertEqual(workbook["Step Outputs"].max_row, 6)
        self.assertEqual(workbook["Tool Calls"].max_row, 4)

        interaction_headers = [
            cell.value
            for cell in workbook["Interactions"][1]
        ]
        tool_headers = [
            cell.value
            for cell in workbook["Tool Calls"][1]
        ]
        self.assertIn("Tool Call Count", interaction_headers)
        self.assertIn("Invalid Source IDs", interaction_headers)
        self.assertIn("Result Preview", tool_headers)
        self.assertEqual(
            workbook["Tool Calls"].freeze_panes,
            "A2",
        )
        self.assertEqual(
            workbook["Tool Calls"]["A1"].fill.fgColor.rgb,
            "001F4E78",
        )
        self.assertFalse(
            workbook["Tool Calls"].sheet_view.showGridLines
        )

        workbook.close()

    def test_failure_records_final_status(self) -> None:
        reporter = ExperimentReporter(
            db_path=str(self.db_path),
            xlsx_path=str(self.xlsx_path),
        )
        experiment_run_id = reporter.start_run(self.config())

        reporter.record_failure(
            experiment_run_id=experiment_run_id,
            question="Failed question?",
            error=RuntimeError("workflow failed"),
        )

        with sqlite3.connect(self.db_path) as connection:
            row = connection.execute(
                """
                SELECT status, final_status, tool_call_count
                FROM interactions
                """
            ).fetchone()

        self.assertEqual(row, ("failed", "failed", 0))

    def test_benchmark_criterion_records_failure_reasons(self) -> None:
        reporter = ExperimentReporter(
            db_path=str(self.db_path),
            xlsx_path=str(self.xlsx_path),
        )
        experiment_run_id = reporter.start_run(self.config())
        response = self.workflow_response()
        response.status = "completed"
        response.step_results = [
            step
            for step in response.step_results
            if step.step_name != "collect_official_sources"
        ]

        reporter.record_workflow_result(
            experiment_run_id=experiment_run_id,
            question="Benchmark question?",
            workflow_response=response,
            benchmark_context={
                "scenario_id": "no-sources",
                "scenario_type": "independent",
                "turn_id": "question",
                "turn_number": 1,
            },
        )

        with sqlite3.connect(self.db_path) as connection:
            row = connection.execute(
                """
                SELECT
                    benchmark_pass,
                    benchmark_failure_reasons
                FROM interactions
                """
            ).fetchone()

        self.assertEqual(row[0], 0)
        self.assertEqual(
            json.loads(row[1]),
            ["no_sources", "invalid_source_ids"],
        )

    def test_migrates_existing_phase_5_database(self) -> None:
        with sqlite3.connect(self.db_path) as connection:
            connection.executescript(
                """
                CREATE TABLE experiment_runs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    run_number INTEGER NOT NULL UNIQUE,
                    started_at_utc TEXT NOT NULL,
                    finished_at_utc TEXT,
                    model_name TEXT NOT NULL,
                    configuration_json TEXT NOT NULL,
                    notes TEXT NOT NULL DEFAULT ''
                );

                CREATE TABLE interactions (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    experiment_run_id INTEGER NOT NULL,
                    question_number INTEGER NOT NULL,
                    workflow_run_id TEXT,
                    session_id TEXT,
                    created_at_utc TEXT NOT NULL,
                    question TEXT NOT NULL,
                    answer TEXT NOT NULL,
                    status TEXT NOT NULL,
                    duration_seconds REAL,
                    input_tokens INTEGER,
                    output_tokens INTEGER,
                    total_tokens INTEGER,
                    time_to_first_token REAL,
                    sources_json TEXT NOT NULL DEFAULT '[]',
                    step_outputs_json TEXT NOT NULL DEFAULT '{}',
                    metrics_json TEXT NOT NULL DEFAULT '{}',
                    error TEXT,
                    UNIQUE (experiment_run_id, question_number)
                );

                INSERT INTO experiment_runs (
                    run_number,
                    started_at_utc,
                    model_name,
                    configuration_json
                )
                VALUES (1, '2026-07-26T00:00:00+00:00', 'old', '{}');

                INSERT INTO interactions (
                    experiment_run_id,
                    question_number,
                    created_at_utc,
                    question,
                    answer,
                    status
                )
                VALUES (
                    1,
                    1,
                    '2026-07-26T00:00:00+00:00',
                    'old question',
                    'old answer',
                    'completed'
                );
                """
            )

        ExperimentReporter(
            db_path=str(self.db_path),
            xlsx_path=str(self.xlsx_path),
        )

        with sqlite3.connect(self.db_path) as connection:
            columns = {
                row[1]
                for row in connection.execute(
                    "PRAGMA table_info(interactions)"
                )
            }
            migrated = connection.execute(
                """
                SELECT
                    tool_call_count,
                    tool_names,
                    final_status
                FROM interactions
                """
            ).fetchone()
            tool_calls_table = connection.execute(
                """
                SELECT name
                FROM sqlite_master
                WHERE type = 'table' AND name = 'tool_calls'
                """
            ).fetchone()

        self.assertTrue(
            {
                "tool_call_count",
                "citation_count",
                "validator_success",
                "final_status",
                "benchmark_pass",
            }.issubset(columns)
        )
        self.assertEqual(migrated, (0, "[]", "unknown"))
        self.assertEqual(tool_calls_table, ("tool_calls",))


if __name__ == "__main__":
    unittest.main()
